import os
import faiss
import pickle
import numpy as np

from sentence_transformers import SentenceTransformer
from rank_bm25 import BM25Okapi

from docx import Document
import openpyxl

DATA_DIR = "data"
INDEX_DIR = "index"

model = SentenceTransformer("BAAI/bge-small-en")

documents = []
metadatas = []

# ------------------------
# ファイル探索（再帰）
# ------------------------
def list_files(root):
    for dirpath, _, filenames in os.walk(root):
        for f in filenames:
            yield os.path.join(dirpath, f)

# ------------------------
# 各ファイル読み込み
# ------------------------

def read_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def read_code(path):
    return read_txt(path)

def read_docx(path):
    doc = Document(path)
    texts = []
    for para in doc.paragraphs:
        texts.append(para.text)
    return "\n".join(texts)

def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    texts = []
    for sheet in wb.worksheets:
        texts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) for cell in row if cell is not None])
            if row_text:
                texts.append(row_text)
    return "\n".join(texts)

# ------------------------
# チャンク分割（タイプ別）
# ------------------------

def chunk_text(text, size=500, overlap=50):
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end])
        start += size - overlap
    return chunks

def chunk_code(text):
    # 関数っぽい単位で分割（簡易）
    lines = text.split("\n")
    chunks = []
    current = []

    for line in lines:
        if line.strip().startswith(("void ", "int ", "class ", "struct ")):
            if current:
                chunks.append("\n".join(current))
                current = []
        current.append(line)

    if current:
        chunks.append("\n".join(current))

    return chunks

def chunk_excel(text):
    # Excelはそのまま小さめチャンク
    return chunk_text(text, size=300, overlap=50)

# ------------------------
# メイン ingest
# ------------------------

def ingest():
    global documents, metadatas

    documents.clear()
    metadatas.clear()

    files = list(list_files(DATA_DIR))
    print(f"Found {len(files)} files")

    for path in files:
        ext = os.path.splitext(path)[1].lower()
        fname = os.path.relpath(path, DATA_DIR)

        try:
            if ext in [".txt", ".md"]:
                text = read_txt(path)
                chunks = chunk_text(text)

            elif ext in [".c", ".cpp", ".h", ".hpp"]:
                text = read_code(path)
                chunks = chunk_code(text)

            elif ext == ".docx":
                text = read_docx(path)
                chunks = chunk_text(text)

            elif ext in [".xlsx", ".xlsm"]:
                text = read_excel(path)
                chunks = chunk_excel(text)

            else:
                continue

            for i, chunk in enumerate(chunks):
                if not chunk.strip():
                    continue

                documents.append(chunk)
                metadatas.append({
                    "source": fname,
                    "chunk_id": i,
                    "type": ext
                })

        except Exception as e:
            print(f"Error processing {path}: {e}")

    print(f"Total chunks: {len(documents)}")

    # ------------------------
    # embedding
    # ------------------------
    embeddings = model.encode(documents, show_progress_bar=True)
    embeddings = np.array(embeddings).astype("float32")

    # ------------------------
    # FAISS
    # ------------------------
    dim = embeddings.shape[1]
    index = faiss.IndexFlatL2(dim)
    index.add(embeddings)

    faiss.write_index(index, f"{INDEX_DIR}/faiss.index")

    # ------------------------
    # metadata保存
    # ------------------------
    with open(f"{INDEX_DIR}/meta.pkl", "wb") as f:
        pickle.dump((documents, metadatas), f)

    # ------------------------
    # BM25
    # ------------------------
    tokenized = [doc.split() for doc in documents]
    bm25 = BM25Okapi(tokenized)

    with open(f"{INDEX_DIR}/bm25.pkl", "wb") as f:
        pickle.dump(bm25, f)

    print("Index build complete")

if __name__ == "__main__":
    ingest()